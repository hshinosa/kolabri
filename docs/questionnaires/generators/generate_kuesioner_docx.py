#!/usr/bin/env python3
"""
Generate professional Questionnaire Word document for Kolabri (Mahasiswa)
"""

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE


def create_questionnaire_document():
    doc = Document()

    # Set default font
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Calibri"
    font.size = Pt(11)

    # Title
    title = doc.add_heading("Kuesioner Kepuasan Pengguna", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = doc.add_paragraph("Kolabri - Aplikasi Pembelajaran Kolaboratif")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.size = Pt(14)
    subtitle.runs[0].font.bold = True

    target = doc.add_paragraph("Target User: Mahasiswa")
    target.alignment = WD_ALIGN_PARAGRAPH.CENTER
    target.runs[0].font.size = Pt(12)

    doc.add_paragraph()

    # Respondent Information
    doc.add_heading("Informasi Responden", 1)
    info_table = doc.add_table(rows=6, cols=2)
    info_table.style = "Light Grid Accent 1"

    info_data = [
        ("Nama (opsional)", "____________"),
        ("NIM", "____________"),
        ("Program Studi", "____________"),
        ("Semester", "____________"),
        ("Tanggal Pengisian", "____________"),
        ("Email (opsional)", "____________"),
    ]

    for i, (label, value) in enumerate(info_data):
        info_table.rows[i].cells[0].text = label
        info_table.rows[i].cells[1].text = value
        info_table.rows[i].cells[0].paragraphs[0].runs[0].font.bold = True

    doc.add_paragraph()

    # Instructions
    doc.add_heading("Petunjuk Pengisian", 1)
    doc.add_paragraph(
        "Berikan penilaian Anda pada setiap pernyataan menggunakan skala berikut:"
    )

    scale_table = doc.add_table(rows=6, cols=2)
    scale_table.style = "Light Grid Accent 1"

    scale_table.rows[0].cells[0].text = "Skor"
    scale_table.rows[0].cells[1].text = "Keterangan"
    scale_table.rows[0].cells[0].paragraphs[0].runs[0].font.bold = True
    scale_table.rows[0].cells[1].paragraphs[0].runs[0].font.bold = True
    scale_table.rows[0].cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    scale_data = [
        ("1", "Sangat Tidak Setuju / Sangat Buruk"),
        ("2", "Tidak Setuju / Buruk"),
        ("3", "Netral / Cukup"),
        ("4", "Setuju / Baik"),
        ("5", "Sangat Setuju / Sangat Baik"),
    ]

    for i, (score, desc) in enumerate(scale_data, 1):
        scale_table.rows[i].cells[0].text = score
        scale_table.rows[i].cells[1].text = desc
        scale_table.rows[i].cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()
    doc.add_page_break()

    # Questionnaire sections
    sections = [
        (
            "A. Kemudahan Penggunaan (Usability)",
            [
                "Kolabri mudah dipelajari dan digunakan",
                "Navigasi antar halaman di Kolabri jelas dan intuitif",
                "Saya dapat menemukan fitur yang saya butuhkan dengan mudah",
                "Tampilan Kolabri menarik dan nyaman dilihat",
                "Teks dan tombol di Kolabri mudah dibaca",
                "Breadcrumb membantu saya mengetahui posisi saya di aplikasi",
                "Pesan error yang ditampilkan mudah dipahami",
            ],
        ),
        (
            "B. Fitur Kelas (Courses)",
            [
                "Saya dapat melihat daftar kelas yang saya ikuti dengan mudah",
                "Proses bergabung dengan kelas menggunakan kode mudah dilakukan",
                "Informasi detail kelas (grup, sesi diskusi, materi) ditampilkan dengan lengkap",
                "Daftar minggu perkuliahan membantu saya mengikuti jadwal kuliah",
                "Saya dapat mengakses materi perkuliahan dengan mudah",
            ],
        ),
        (
            "C. Fitur Kelompok (Groups)",
            [
                "Proses membuat kelompok baru mudah dilakukan",
                "Kode undangan kelompok mudah dibagikan ke anggota lain",
                "Saya dapat melihat daftar anggota kelompok dengan jelas",
                "Fitur manajemen anggota (untuk ketua kelompok) berfungsi dengan baik",
                "Log aktivitas kelompok membantu saya memantau perubahan",
            ],
        ),
        (
            "D. Fitur Ruang Diskusi (Chat Spaces)",
            [
                "Chat room loading dengan cepat",
                "Pesan yang saya kirim muncul secara real-time",
                "Fitur upload file di chat berfungsi dengan baik",
                "Fitur edit dan hapus pesan membantu saya memperbaiki kesalahan",
                "Fitur pin pesan membantu saya menandai pesan penting",
                "Fitur search pesan membantu saya menemukan informasi lama",
                "Indikator status koneksi (connected/reconnecting) jelas",
                "Fitur tutup sesi diskusi berguna untuk mengakhiri diskusi",
                "Summary sesi diskusi membantu saya mengingat poin penting",
            ],
        ),
        (
            "E. Fitur Pre-read (Materi Sebelum Sesi)",
            [
                "Saya dapat mengakses materi pre-read dengan mudah",
                "Viewer PDF/Video berfungsi dengan baik",
                "Fitur menandai pre-read selesai membantu saya tracking progress",
                "Pre-read membantu saya mempersiapkan diri sebelum diskusi",
            ],
        ),
        (
            "F. Fitur Refleksi (Reflections)",
            [
                "Saya dapat membuat refleksi dengan mudah",
                "Template refleksi membantu saya menulis refleksi yang terstruktur",
                "Fitur tag membantu saya mengorganisir refleksi",
                "Fitur search dan filter refleksi berfungsi dengan baik",
                "Analytics refleksi memberikan insight yang berguna tentang kebiasaan refleksi saya",
                "Refleksi membantu saya mengevaluasi proses belajar saya",
            ],
        ),
        (
            "G. Fitur AI Chat",
            [
                "AI Chat mudah diakses dan digunakan",
                "Response AI muncul dengan cepat (streaming)",
                "Jawaban AI relevan dan membantu pembelajaran saya",
                "Response AI dengan formatting (list, code, table) mudah dibaca",
                "Fitur riwayat chat membantu saya mengakses percakapan lama",
                "Fitur simpan materi dari AI response berguna untuk referensi",
                "Fitur search di AI Chat membantu menemukan informasi lama",
                "Sitasi materi yang diberikan AI akurat dan relevan",
            ],
        ),
        (
            "H. Fitur Goals (Tujuan Pembelajaran)",
            [
                "Fitur goals membantu saya menetapkan tujuan pembelajaran",
                "Saya dapat membuat dan mengedit goals dengan mudah",
                "Goals membantu saya fokus pada tujuan diskusi",
            ],
        ),
        (
            "I. Fitur Profil & Pengaturan",
            [
                "Saya dapat mengedit profil dengan mudah",
                "Fitur upload avatar berfungsi dengan baik",
                "Statistik profil memberikan informasi yang berguna",
                "Pengaturan preferensi mudah diakses dan diubah",
            ],
        ),
        (
            "J. Kinerja Sistem (Performance)",
            [
                "Kolabri loading dengan cepat",
                "Kolabri responsif saat digunakan",
                "Saya jarang mengalami error atau crash",
                "Kolabri dapat diakses dengan baik di perangkat saya",
            ],
        ),
        (
            "K. Kepuasan Keseluruhan",
            [
                "Secara keseluruhan, saya puas dengan Kolabri",
                "Kolabri membantu proses belajar saya",
                "Kolabri memudahkan kolaborasi dengan teman sekelompok",
                "Saya akan merekomendasikan Kolabri kepada teman",
                "Saya akan menggunakan Kolabri lagi di masa depan",
            ],
        ),
    ]

    # Create questionnaire tables
    for section_title, statements in sections:
        doc.add_heading(f"Bagian {section_title}", 1)

        table = doc.add_table(rows=1, cols=7)
        table.style = "Light Grid Accent 1"

        # Header
        headers = ["No", "Pernyataan", "1", "2", "3", "4", "5"]
        for i, header in enumerate(headers):
            table.rows[0].cells[i].text = header
            table.rows[0].cells[i].paragraphs[0].runs[0].font.bold = True
            table.rows[0].cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Set column widths
        widths = [
            Inches(0.4),
            Inches(3.5),
            Inches(0.4),
            Inches(0.4),
            Inches(0.4),
            Inches(0.4),
            Inches(0.4),
        ]
        for i, width in enumerate(widths):
            for cell in table.columns[i].cells:
                cell.width = width

        # Add statements
        for i, statement in enumerate(statements, 1):
            row = table.add_row()
            row.cells[0].text = f"{section_title[0]}{i}"
            row.cells[1].text = statement

            for j in range(2, 7):
                row.cells[j].text = "☐"
                row.cells[j].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

            row.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

        doc.add_paragraph()

    doc.add_page_break()

    # Open-ended questions
    doc.add_heading("Bagian L: Pertanyaan Terbuka", 1)

    open_questions = [
        "L1. Fitur apa yang paling Anda sukai di Kolabri? Mengapa?",
        "L2. Fitur apa yang paling kurang Anda sukai atau perlu diperbaiki? Mengapa?",
        "L3. Apakah ada fitur yang Anda harapkan ada di Kolabri tetapi belum tersedia?",
        "L4. Saran atau masukan lain untuk pengembangan Kolabri?",
    ]

    for question in open_questions:
        p = doc.add_paragraph()
        p.add_run(question).font.bold = True

        # Add answer space
        for _ in range(3):
            doc.add_paragraph("_" * 80)

        doc.add_paragraph()

    doc.add_page_break()

    # Summary scoring
    doc.add_heading("Ringkasan Penilaian", 1)
    doc.add_paragraph("(Diisi oleh peneliti/admin)")

    summary_table = doc.add_table(rows=1, cols=4)
    summary_table.style = "Light Grid Accent 1"

    summary_headers = ["Bagian", "Total Skor", "Skor Maksimum", "Persentase"]
    for i, header in enumerate(summary_headers):
        summary_table.rows[0].cells[i].text = header
        summary_table.rows[0].cells[i].paragraphs[0].runs[0].font.bold = True
        summary_table.rows[0].cells[i].paragraphs[
            0
        ].alignment = WD_ALIGN_PARAGRAPH.CENTER

    summary_data = [
        ("A. Kemudahan Penggunaan", "35"),
        ("B. Fitur Kelas", "25"),
        ("C. Fitur Kelompok", "25"),
        ("D. Fitur Ruang Diskusi", "45"),
        ("E. Fitur Pre-read", "20"),
        ("F. Fitur Refleksi", "30"),
        ("G. Fitur AI Chat", "40"),
        ("H. Fitur Goals", "15"),
        ("I. Fitur Profil", "20"),
        ("J. Kinerja Sistem", "20"),
        ("K. Kepuasan Keseluruhan", "25"),
    ]

    for section, max_score in summary_data:
        row = summary_table.add_row()
        row.cells[0].text = section
        row.cells[1].text = ""
        row.cells[2].text = max_score
        row.cells[3].text = ""
        for i in range(1, 4):
            row.cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Total row
    total_row = summary_table.add_row()
    total_row.cells[0].text = "TOTAL"
    total_row.cells[0].paragraphs[0].runs[0].font.bold = True
    total_row.cells[1].text = ""
    total_row.cells[2].text = "300"
    total_row.cells[2].paragraphs[0].runs[0].font.bold = True
    total_row.cells[3].text = ""
    for i in range(1, 4):
        total_row.cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()

    # Interpretation
    doc.add_heading("Interpretasi Skor", 1)

    interp_table = doc.add_table(rows=6, cols=2)
    interp_table.style = "Light Grid Accent 1"

    interp_table.rows[0].cells[0].text = "Persentase"
    interp_table.rows[0].cells[1].text = "Kategori"
    interp_table.rows[0].cells[0].paragraphs[0].runs[0].font.bold = True
    interp_table.rows[0].cells[1].paragraphs[0].runs[0].font.bold = True

    interp_data = [
        ("85-100%", "Sangat Baik"),
        ("70-84%", "Baik"),
        ("55-69%", "Cukup"),
        ("40-54%", "Kurang"),
        ("<40%", "Sangat Kurang"),
    ]

    for i, (percentage, category) in enumerate(interp_data, 1):
        interp_table.rows[i].cells[0].text = percentage
        interp_table.rows[i].cells[1].text = category
        interp_table.rows[i].cells[0].paragraphs[
            0
        ].alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()
    doc.add_paragraph()

    # Thank you
    thanks = doc.add_paragraph("Terima kasih atas partisipasi Anda!")
    thanks.alignment = WD_ALIGN_PARAGRAPH.CENTER
    thanks.runs[0].font.size = Pt(14)
    thanks.runs[0].font.bold = True
    thanks.runs[0].font.color.rgb = RGBColor(0, 102, 204)

    # Save document
    output_path = "/Users/hshino/Kuliah/ProjectTA/docs/Kuesioner_Mahasiswa_Kolabri.docx"
    doc.save(output_path)
    print(f"✓ Questionnaire document saved to: {output_path}")


if __name__ == "__main__":
    create_questionnaire_document()
