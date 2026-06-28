#!/usr/bin/env python3
"""
Generate compact Questionnaire Word document for Kolabri (Mahasiswa)
~32 questions total
"""

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_compact_questionnaire():
    doc = Document()
    
    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)
    
    # Title
    title = doc.add_heading('Kuesioner Kepuasan Pengguna', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    subtitle = doc.add_paragraph('Kolabri - Aplikasi Pembelajaran Kolaboratif')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.size = Pt(14)
    subtitle.runs[0].font.bold = True
    
    target = doc.add_paragraph('Target User: Mahasiswa')
    target.alignment = WD_ALIGN_PARAGRAPH.CENTER
    target.runs[0].font.size = Pt(12)
    
    doc.add_paragraph()
    
    # Respondent Information
    doc.add_heading('Informasi Responden', 1)
    info_table = doc.add_table(rows=5, cols=2)
    info_table.style = 'Light Grid Accent 1'
    
    info_data = [
        ('Nama (opsional)', '____________'),
        ('NIM', '____________'),
        ('Program Studi', '____________'),
        ('Semester', '____________'),
        ('Tanggal Pengisian', '____________')
    ]
    
    for i, (label, value) in enumerate(info_data):
        info_table.rows[i].cells[0].text = label
        info_table.rows[i].cells[1].text = value
        info_table.rows[i].cells[0].paragraphs[0].runs[0].font.bold = True
    
    doc.add_paragraph()
    
    # Instructions
    doc.add_heading('Petunjuk Pengisian', 1)
    doc.add_paragraph('Berikan penilaian Anda pada setiap pernyataan menggunakan skala berikut:')
    
    scale_table = doc.add_table(rows=6, cols=2)
    scale_table.style = 'Light Grid Accent 1'
    
    scale_table.rows[0].cells[0].text = 'Skor'
    scale_table.rows[0].cells[1].text = 'Keterangan'
    scale_table.rows[0].cells[0].paragraphs[0].runs[0].font.bold = True
    scale_table.rows[0].cells[1].paragraphs[0].runs[0].font.bold = True
    scale_table.rows[0].cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    scale_data = [
        ('1', 'Sangat Tidak Setuju'),
        ('2', 'Tidak Setuju'),
        ('3', 'Netral'),
        ('4', 'Setuju'),
        ('5', 'Sangat Setuju')
    ]
    
    for i, (score, desc) in enumerate(scale_data, 1):
        scale_table.rows[i].cells[0].text = score
        scale_table.rows[i].cells[1].text = desc
        scale_table.rows[i].cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph()
    doc.add_page_break()
    
    # Compact questionnaire sections (~32 questions)
    sections = [
        ('A. Kemudahan Penggunaan', [
            'Kolabri mudah dipelajari dan digunakan',
            'Navigasi antar halaman jelas dan intuitif',
            'Tampilan visual menarik dan nyaman dilihat',
            'Pesan error mudah dipahami'
        ]),
        ('B. Fitur Kelas & Materi', [
            'Saya dapat melihat daftar kelas yang diikuti dengan mudah',
            'Proses bergabung dengan kelas menggunakan kode mudah dilakukan',
            'Informasi detail kelas ditampilkan dengan lengkap',
            'Saya dapat mengakses materi perkuliahan dengan mudah',
            'Materi pre-read membantu saya mempersiapkan diskusi'
        ]),
        ('C. Fitur Kolaborasi', [
            'Proses membuat/bergabung dengan kelompok mudah dilakukan',
            'Chat room berfungsi dengan baik (real-time, upload file)',
            'Fitur edit, hapus, dan pin pesan berguna untuk diskusi',
            'Fitur search pesan membantu menemukan informasi',
            'Summary sesi diskusi membantu mengingat poin penting'
        ]),
        ('D. Fitur Refleksi & Pembelajaran', [
            'Saya dapat membuat refleksi dengan mudah',
            'Template refleksi membantu menulis yang terstruktur',
            'Analytics refleksi memberikan insight yang berguna',
            'Fitur goals membantu menetapkan tujuan pembelajaran',
            'Refleksi membantu saya mengevaluasi proses belajar'
        ]),
        ('E. Fitur AI Chat', [
            'AI Chat mudah diakses dan digunakan',
            'Response AI muncul dengan cepat dan relevan',
            'Fitur riwayat chat dan saved materials berguna',
            'Sitasi materi dari AI akurat dan membantu'
        ]),
        ('F. Kinerja & Keandalan', [
            'Kolabri loading dengan cepat',
            'Kolabri responsif dan jarang error/crash',
            'Kolabri dapat diakses dengan baik di perangkat saya'
        ]),
        ('G. Kepuasan Keseluruhan', [
            'Secara keseluruhan, saya puas dengan Kolabri',
            'Kolabri membantu proses belajar saya',
            'Kolabri memudahkan kolaborasi dengan teman',
            'Saya akan merekomendasikan Kolabri kepada teman'
        ])
    ]
    
    # Create questionnaire tables
    question_count = 0
    for section_title, statements in sections:
        doc.add_heading(f'Bagian {section_title}', 1)
        
        table = doc.add_table(rows=1, cols=7)
        table.style = 'Light Grid Accent 1'
        
        # Header
        headers = ['No', 'Pernyataan', '1', '2', '3', '4', '5']
        for i, header in enumerate(headers):
            table.rows[0].cells[i].text = header
            table.rows[0].cells[i].paragraphs[0].runs[0].font.bold = True
            table.rows[0].cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Set column widths
        widths = [Inches(0.4), Inches(3.5), Inches(0.4), Inches(0.4), Inches(0.4), Inches(0.4), Inches(0.4)]
        for i, width in enumerate(widths):
            for cell in table.columns[i].cells:
                cell.width = width
        
        # Add statements
        for i, statement in enumerate(statements, 1):
            question_count += 1
            row = table.add_row()
            row.cells[0].text = f'{section_title[0]}{i}'
            row.cells[1].text = statement
            
            for j in range(2, 7):
                row.cells[j].text = '☐'
                row.cells[j].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            row.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        doc.add_paragraph()
    
    doc.add_page_break()
    
    # Open-ended questions (3 questions)
    doc.add_heading('Bagian H: Pertanyaan Terbuka', 1)
    
    open_questions = [
        'H1. Fitur apa yang paling Anda sukai di Kolabri?',
        'H2. Apa yang perlu diperbaiki atau ditambahkan?',
        'H3. Saran lain untuk pengembangan Kolabri?'
    ]
    
    for question in open_questions:
        p = doc.add_paragraph()
        p.add_run(question).font.bold = True
        
        # Add answer space
        for _ in range(2):
            doc.add_paragraph('_' * 80)
        
        doc.add_paragraph()
    
    doc.add_page_break()
    
    # Summary scoring
    doc.add_heading('Ringkasan Penilaian', 1)
    doc.add_paragraph('(Diisi oleh peneliti/admin)')
    
    summary_table = doc.add_table(rows=1, cols=4)
    summary_table.style = 'Light Grid Accent 1'
    
    summary_headers = ['Bagian', 'Total Skor', 'Skor Maksimum', 'Persentase']
    for i, header in enumerate(summary_headers):
        summary_table.rows[0].cells[i].text = header
        summary_table.rows[0].cells[i].paragraphs[0].runs[0].font.bold = True
        summary_table.rows[0].cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    summary_data = [
        ('A. Kemudahan Penggunaan', '20'),
        ('B. Fitur Kelas & Materi', '25'),
        ('C. Fitur Kolaborasi', '25'),
        ('D. Fitur Refleksi & Pembelajaran', '25'),
        ('E. Fitur AI Chat', '20'),
        ('F. Kinerja & Keandalan', '15'),
        ('G. Kepuasan Keseluruhan', '20')
    ]
    
    for section, max_score in summary_data:
        row = summary_table.add_row()
        row.cells[0].text = section
        row.cells[1].text = ''
        row.cells[2].text = max_score
        row.cells[3].text = ''
        for i in range(1, 4):
            row.cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Total row
    total_row = summary_table.add_row()
    total_row.cells[0].text = 'TOTAL'
    total_row.cells[0].paragraphs[0].runs[0].font.bold = True
    total_row.cells[1].text = ''
    total_row.cells[2].text = '150'
    total_row.cells[2].paragraphs[0].runs[0].font.bold = True
    total_row.cells[3].text = ''
    for i in range(1, 4):
        total_row.cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph()
    
    # Interpretation
    doc.add_heading('Interpretasi Skor', 1)
    
    interp_table = doc.add_table(rows=6, cols=2)
    interp_table.style = 'Light Grid Accent 1'
    
    interp_table.rows[0].cells[0].text = 'Persentase'
    interp_table.rows[0].cells[1].text = 'Kategori'
    interp_table.rows[0].cells[0].paragraphs[0].runs[0].font.bold = True
    interp_table.rows[0].cells[1].paragraphs[0].runs[0].font.bold = True
    
    interp_data = [
        ('85-100%', 'Sangat Baik'),
        ('70-84%', 'Baik'),
        ('55-69%', 'Cukup'),
        ('40-54%', 'Kurang'),
        ('<40%', 'Sangat Kurang')
    ]
    
    for i, (percentage, category) in enumerate(interp_data, 1):
        interp_table.rows[i].cells[0].text = percentage
        interp_table.rows[i].cells[1].text = category
        interp_table.rows[i].cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph()
    doc.add_paragraph()
    
    # Thank you
    thanks = doc.add_paragraph('Terima kasih atas partisipasi Anda!')
    thanks.alignment = WD_ALIGN_PARAGRAPH.CENTER
    thanks.runs[0].font.size = Pt(14)
    thanks.runs[0].font.bold = True
    thanks.runs[0].font.color.rgb = RGBColor(0, 102, 204)
    
    # Save document
    output_path = '/Users/hshino/Kuliah/ProjectTA/docs/Kuesioner_Mahasiswa_Kolabri_Compact.docx'
    doc.save(output_path)
    print(f"✓ Compact questionnaire saved to: {output_path}")
    print(f"  Total questions: {question_count} (reduced from 60)")

if __name__ == '__main__':
    create_compact_questionnaire()
