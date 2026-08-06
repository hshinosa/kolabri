#!/usr/bin/env python3
"""Extract text structure from the original D3 Capstone docx."""

from __future__ import annotations

from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parent
SRC = ROOT / "[D3] Capstone Project Final Report and Documentation.docx"
OUT = ROOT / "extracted_original.md"


def main() -> None:
    doc = Document(str(SRC))
    lines: list[str] = []
    for p in doc.paragraphs:
        style = p.style.name if p.style else "Normal"
        text = p.text or ""
        if style.startswith("Heading"):
            level = "".join(ch for ch in style if ch.isdigit()) or "1"
            lines.append(f"{'#' * int(level)} {text}")
        else:
            lines.append(text)
    OUT.write_text("\n".join(lines), encoding="utf-8")
    print(f"Extracted {len(doc.paragraphs)} paragraphs -> {OUT}")
    print(f"Tables: {len(doc.tables)}")


if __name__ == "__main__":
    main()
