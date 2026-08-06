#!/usr/bin/env python3
"""
Compile D3 Capstone Final Report for Kolabri from:
1) Original D3 template content that is already filled (cover + background)
2) Actual repository implementation under ProjectTA (code-first)

No invented metrics. Claims are limited to structures, modules, and flows
that exist in the codebase at compile time.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "[D3] Capstone Project Final Report and Documentation - Kolabri.docx"


def set_run_font(run, *, name: str = "Times New Roman", size: int = 12, bold: bool = False, italic: bool = False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def add_para(
    doc: Document,
    text: str = "",
    *,
    style: str | None = None,
    align=None,
    bold: bool = False,
    italic: bool = False,
    size: int = 12,
    space_after: float = 6,
    space_before: float = 0,
    first_line_indent: float | None = None,
):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    if first_line_indent is not None:
        pf.first_line_indent = Cm(first_line_indent)
    if text:
        run = p.add_run(text)
        set_run_font(run, size=size, bold=bold, italic=italic)
    return p


def add_heading(doc: Document, text: str, level: int = 1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, size=14 if level == 1 else 12, bold=True)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(8)
    return p


def add_bullets(doc: Document, items: list[str], *, size: int = 12):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        run = p.add_run(item)
        set_run_font(run, size=size)


def add_numbered(doc: Document, items: list[str], *, size: int = 12, start: int = 1):
    for i, item in enumerate(items):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        p.paragraph_format.left_indent = Cm(1.0)
        p.paragraph_format.first_line_indent = Cm(-0.5)
        run = p.add_run(f"{start + i}. {item}")
        set_run_font(run, size=size)


def add_table(doc: Document, headers: list[str], rows: list[list[str]]):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(h)
        set_run_font(run, size=11, bold=True)
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = table.rows[r_i + 1].cells[c_i]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(val)
            set_run_font(run, size=10)
    doc.add_paragraph()
    return table


def build() -> Path:
    doc = Document()

    # Page setup A4
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.5)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    # ========== COVER (from original D3) ==========
    add_para(doc, "Capstone Project", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=16, space_after=4)
    add_para(doc, "Final Report and Documentation", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=16, space_after=18)
    add_para(doc, "Web Chatbot Kolaboratif", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=14, space_after=2)
    add_para(doc, "(Kolabri)", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=14, space_after=18)
    add_para(doc, "Team Members:", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=12, space_after=6)

    members = [
        "Soraya Haidar Salma (1302223006) — System Analyst",
        "Irham Baehaqi (1302220063) — Backend Developer (Chat & User Management)",
        "Muhammad Hashfi Hadyan (1302220079) — Backend Developer (AI Integration & Analytics)",
        "Mochammad Rizky Septian (1302220121) — Frontend Developer (Chat UI & Group Space)",
        "Ahmad Fadli Akbar (1302220126) — Frontend Developer (Dashboard & Visualization)",
        "Benedict Arvin Indra Puteprasa (1302223136) — QA Engineer",
    ]
    for i, m in enumerate(members, 1):
        add_para(doc, f"{i}. {m}", align=WD_ALIGN_PARAGRAPH.CENTER, size=11, space_after=2)

    add_para(doc, "", space_after=12)
    add_para(doc, "Program Studi Sarjana Rekayasa Perangkat Lunak", align=WD_ALIGN_PARAGRAPH.CENTER, size=12, space_after=2)
    add_para(doc, "Fakultas Informatika", align=WD_ALIGN_PARAGRAPH.CENTER, size=12, space_after=2)
    add_para(doc, "Universitas Telkom", align=WD_ALIGN_PARAGRAPH.CENTER, size=12, space_after=2)
    add_para(doc, "2026", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=12, space_after=12)

    # ========== ABSTRAK ==========
    add_heading(doc, "Abstrak", level=1)
    abstrak = (
        "Kolabri (Web Chatbot Kolaboratif) adalah platform pembelajaran kolaboratif berbasis web "
        "untuk perguruan tinggi yang menghubungkan mahasiswa, dosen, dan agen AI dalam satu alur "
        "diskusi terstruktur. Sistem diimplementasikan sebagai tiga layanan: Client App (Laravel 12 "
        "BFF + Inertia.js/React TypeScript, port 8000), Core API (Node.js/Express TypeScript, "
        "Prisma/PostgreSQL, MongoDB, Socket.IO, Redis, port 3000), dan AI Engine (Python FastAPI, "
        "port 8001) dengan Qdrant sebagai vector store. Data domain (user, course, group, session, "
        "goal, reflection, week/material, AI provider) disimpan utama di PostgreSQL; MongoDB "
        "dipakai untuk log runtime chat/aktivitas; Qdrant untuk embedding RAG. Alur mahasiswa "
        "mengikuti siklus regulasi belajar: pre-read materi minggu → penetapan learning goal "
        "(validasi kata kerja Bloom + umpan balik AI) → diskusi kelompok real-time dengan AI "
        "orchestration (RAG, sitasi, scaffolding) → penutupan sesi → refleksi (session/weekly). "
        "Logic Listener dan pipeline intervensi memantau off-topic (kemiripan embedding), silence "
        "(ambang 10 menit di socket gate), dan partisipasi tidak merata (koefisien Gini), dengan "
        "eskalasi bertahap (new → nudge → probe-blocker → flag-lecturer → resolved). Guardrails "
        "memfilter input/output (academic dishonesty, off-topic, toxicity, PII, injection, filter "
        "Socratic). Dosen memperoleh dashboard, analytics kursus/kelompok/mahasiswa, discussion "
        "health, dan knowledge base; admin mengelola pengguna, master data, provider AI, audit "
        "log, dan usage stats. Chat pribadi (AiChat) menyediakan mode 1:1 dengan RAG. AiUsage "
        "(token/latency/biaya), AuditLog, dan ExportJob mendukung monitoring serta analisis "
        "pembelajaran. Laporan ini mendokumentasikan desain, implementasi, evaluasi berbasis "
        "pengujian perangkat lunak, tantangan, dan kesimpulan yang selaras dengan kode sumber "
        "repositori Kolabri."
    )
    add_para(doc, abstrak, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line_indent=1.0, space_after=8)
    add_para(
        doc,
        "Katakunci: pembelajaran kolaboratif; chatbot edukasi; RAG; Self-Regulated Learning; "
        "Logic Listener; guardrails; Socket.IO; FastAPI; Laravel; analytics pembelajaran",
        italic=True,
        space_after=12,
    )

    # ========== BACKGROUND (from original filled D3 — preserved) ==========
    add_heading(doc, "Background, Motivation, and Problem Definition", level=1)
    add_para(doc, "a. Background", bold=True, space_after=6)

    bg_paras = [
        (
            "Pembelajaran kolaboratif telah terbukti secara empiris memberikan dampak positif yang "
            "signifikan dalam ekosistem pendidikan tinggi, terutama dalam meningkatkan keterampilan "
            "berpikir kritis, kemampuan komunikasi, serta retensi pengetahuan mahasiswa. Meskipun "
            "manfaatnya sangat jelas, implementasi strategi ini di lingkungan akademik sering kali "
            "terhambat oleh tantangan sistemik yang mengurangi efektivitasnya. Hambatan ini tidak "
            "hanya bersifat pedagogis, tetapi juga menyentuh aspek teknis, sehingga menuntut adanya "
            "solusi terintegrasi yang mampu menjembatani teori pembelajaran dengan pemanfaatan "
            "teknologi cerdas."
        ),
        (
            "Secara pedagogis, tantangan utama yang sering muncul adalah ketimpangan partisipasi "
            "(participation inequity) dalam diskusi kelompok. Fenomena seperti free-riding atau "
            "dominasi diskusi oleh segelintir anggota menyebabkan distribusi manfaat pembelajaran "
            "yang tidak merata di antara mahasiswa. Selain itu, diskusi sering kali melebar keluar "
            "dari topik (off-topic discussion), yang mengakibatkan inefisiensi waktu dan kegagalan "
            "dalam mencapai tujuan pembelajaran. Situasi ini diperburuk oleh keterbatasan kapasitas "
            "dosen dalam memantau dinamika multipel kelompok secara simultan, khususnya di kelas "
            "besar. Akibatnya, tanpa adanya dokumentasi dan ringkasan otomatis yang memadai, proses "
            "diskusi sering kali berlalu tanpa jejak rekam yang baik, menyulitkan proses refleksi "
            "maupun penilaian objektif."
        ),
        (
            "Di sisi lain, kemajuan pesat Large Language Models (LLM) seperti Google Gemini membuka "
            "peluang baru untuk mengatasi hambatan tersebut melalui pengembangan agen pedagogis "
            "cerdas. Martha et al. (2023) telah membuktikan bahwa integrasi scaffolding adaptif "
            "dalam agen pedagogis dapat secara signifikan meningkatkan keterampilan Self-Regulated "
            "Learning (SRL) dan Co-Regulated Learning (CoRL). Potensi ini menawarkan jalan keluar "
            "bagi masalah monitoring dan fasilitasi diskusi, di mana AI dapat berperan sebagai "
            "fasilitator yang membantu menjaga fokus dan keseimbangan partisipasi dalam kelompok."
        ),
        (
            "Namun, transisi menuju penerapan AI dalam pendidikan memerlukan penanganan tantangan "
            "teknis yang serius. Risiko utama yang dihadapi adalah “halusinasi” LLM, di mana model "
            "rentan menghasilkan informasi yang meyakinkan namun faktualnya keliru. Untuk memitigasi "
            "risiko ini, pendekatan Retrieval-Augmented Generation (RAG) menjadi krusial sebagaimana "
            "ditunjukkan oleh Lewis et al. (2021), karena mampu membatasi respons AI pada konteks "
            "dokumen yang valid. Lebih jauh lagi, pengembangan sistem ini menghadapi kompleksitas "
            "dalam Requirement Engineering, di mana metode Agile tradisional sering kali gagal "
            "menangani ketidakpastian spesifikasi sistem berbasis AI (Hoy & Xu, 2023). Oleh karena "
            "itu, diperlukan arsitektur sistem yang tidak hanya cerdas dan akurat, tetapi juga "
            "mampu menjamin integritas data melalui pembentukan log terstruktur yang kompatibel "
            "dengan standar Educational Process Mining untuk keperluan riset pedagogis yang valid."
        ),
    ]
    for t in bg_paras:
        add_para(doc, t, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line_indent=1.0, space_after=8)

    add_para(doc, "b. Motivation", bold=True, space_after=6)
    add_para(
        doc,
        "Proyek ini termotivasi oleh tiga kesenjangan (gaps) utama antara kebutuhan pedagogis dan "
        "solusi teknologi yang tersedia:",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
        space_after=6,
    )
    add_bullets(
        doc,
        [
            "Gap Teoritis-Praktis: Teori regulasi belajar Zimmerman yang telah mapan secara teoretis "
            "(Forethought–Performance–Reflection) belum terimplementasi secara komprehensif dalam "
            "sistem chatbot edukasi. Sebagian besar chatbot yang ada hanya berfungsi sebagai "
            "question-answering system pasif, bukan sebagai fasilitator aktif yang mendukung "
            "seluruh siklus regulasi belajar.",
            "Gap Teknologi: Sistem chatbot edukasi yang ada mayoritas tidak mengintegrasikan "
            "mekanisme RAG untuk mencegah halusinasi, tidak memiliki Logic Listener untuk deteksi "
            "dinamika kelompok secara real-time, dan tidak menghasilkan data log yang terstruktur "
            "untuk analisis pembelajaran.",
            "Gap Infrastruktur Penelitian: Ketiadaan sistem yang menghasilkan data research-grade "
            "dengan struktur log yang mengikuti taksonomi standar (Gen-SRL) menyulitkan peneliti "
            "pendidikan untuk melakukan analisis kausal tentang efektivitas intervensi pembelajaran.",
        ],
    )

    add_para(doc, "c. Problem Definition", bold=True, space_before=8, space_after=6)
    add_para(
        doc,
        "Berdasarkan analisis latar belakang, proyek ini merumuskan permasalahan utama sebagai berikut:",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
        space_after=6,
    )
    add_para(doc, "Permasalahan Inti:", bold=True, space_after=4)
    add_para(
        doc,
        "Bagaimana merancang dan mengimplementasikan sistem Web Chatbot Kolaboratif berbasis AI yang mampu:",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
        space_after=4,
    )
    add_bullets(
        doc,
        [
            "Memfasilitasi tiga tingkat regulasi belajar (SRL, CoRL, SSRL) sesuai siklus Zimmerman",
            "Mendeteksi dan mengintervensi dinamika kelompok yang tidak produktif (pasivitas, dominasi, off-topic) secara otomatis",
            "Menjamin akurasi respons AI melalui RAG dan Guardrails",
            "Menghasilkan data log terstruktur untuk keperluan monitoring dosen dan riset pedagogis",
        ],
    )
    add_para(doc, "Sub-Permasalahan Spesifik per Role:", bold=True, space_before=6, space_after=4)
    add_bullets(
        doc,
        [
            "System Analyst (Soraya): Bagaimana menerjemahkan kebutuhan pedagogis abstrak (fase regulasi, deteksi off-topic) menjadi spesifikasi UML yang presisi dan implementable?",
            "Backend AI (Hashfi): Bagaimana mengintegrasikan RAG dengan Guardrails berlapis untuk menjamin akurasi dan keamanan, serta merancang arsitektur asinkron yang scalable?",
            "Backend Chat (Irham): Bagaimana membangun infrastruktur komunikasi real-time dengan latensi rendah yang mendukung transactional logging untuk integritas data?",
            "Frontend Chat (Rizky): Bagaimana merancang antarmuka yang mendorong otonomi mahasiswa (SRL) dan memvisualisasikan intervensi SSRL secara intuitif?",
            "Frontend Dashboard (Akbar): Bagaimana menyajikan data analitik kompleks (risiko kelompok, process mining) dalam visualisasi yang actionable bagi dosen?",
            "QA (Arvin): Bagaimana memvalidasi fungsionalitas fitur cerdas (deteksi off-topic, scaffolding adaptif) dan integritas data log untuk riset?",
        ],
    )

    # ========== TEAM ==========
    add_heading(doc, "Team Members and Role", level=1)
    add_para(
        doc,
        "Pembagian peran disusun agar setiap area kritis sistem (spesifikasi, domain chat/user, "
        "AI/analytics, UI diskusi, dashboard, dan QA) memiliki pemilik yang jelas. Justifikasi "
        "didasarkan pada domain teknis repositori Kolabri, bukan pada pembagian formal belaka.",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
        first_line_indent=1.0,
        space_after=8,
    )
    add_table(
        doc,
        ["Nama / NIM", "Peran", "Justifikasi terhadap ruang lingkup sistem"],
        [
            [
                "Soraya Haidar Salma\n1302223006",
                "System Analyst",
                "Menerjemahkan kebutuhan pedagogis (SRL/CoRL/SSRL, pre-read, goal, refleksi, intervensi) menjadi spesifikasi use case, alur, dan kontrak antar-layanan (BFF ↔ Core API ↔ AI Engine).",
            ],
            [
                "Irham Baehaqi\n1302220063",
                "Backend Developer\n(Chat & User Management)",
                "Fokus domain Core API: autentikasi/otorisasi (JWT, role student/lecturer/admin), grup, session discussion, Socket.IO (join_room, send_message, presence, pin/delete), pre-read gate, eskalasi, dan model data Prisma/Mongo.",
            ],
            [
                "Muhammad Hashfi Hadyan\n1302220079",
                "Backend Developer\n(AI Integration & Analytics)",
                "Fokus AI Engine dan integrasi: RAG (Qdrant, embedding, reranker, week context), orchestration, Logic Listener, guardrails, validasi goal, NLP analytics (HOT/lexical), process mining (plan-vs-reality, XES), serta tracking AiUsage.",
            ],
            [
                "Mochammad Rizky Septian\n1302220121",
                "Frontend Developer\n(Chat UI & Group Space)",
                "UI mahasiswa: course/group/session discussion, pre-read, goal, chat room real-time (socket.io-client), ringkasan sesi, chat AI pribadi (AiChat), serta fitur chat (pin, reply, upload, search).",
            ],
            [
                "Ahmad Fadli Akbar\n1302220126",
                "Frontend Developer\n(Dashboard & Visualization)",
                "UI dosen/admin: dashboard, analytics kursus/kelompok/mahasiswa (Chart.js/Recharts), discussion health, knowledge base, attendance, AI settings, user management, audit log.",
            ],
            [
                "Benedict Arvin Indra Puteprasa\n1302223136",
                "QA Engineer",
                "Perencanaan dan pelaksanaan pengujian: unit/integration (core-api *.test.ts, ai-engine tests/, client tests), smoke/UAT alur utama (login multi-role, pre-read→goal→chat→close→reflection), dan cek regresi fitur cerdas.",
            ],
        ],
    )

    # ========== RELATED WORKS ==========
    add_heading(doc, "Related Works", level=1)
    add_para(
        doc,
        "Bagian ini merangkum landasan karya terkait yang menjadi acuan desain Kolabri. "
        "Uraian dibatasi pada referensi yang memang dipakai dalam rumusan masalah dan "
        "implementasi teknis di repositori, tanpa menambahkan klaim empiris di luar cakupan proyek.",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
        first_line_indent=1.0,
        space_after=8,
    )
    add_para(doc, "1. Scaffolding adaptif dan regulasi belajar", bold=True, space_after=4)
    add_para(
        doc,
        "Martha et al. (2023) menunjukkan bahwa agen pedagogis dengan scaffolding adaptif dapat "
        "mendukung keterampilan SRL dan CoRL. Kolabri mengadopsi prinsip ini melalui filter "
        "Socratic (socratic_filter), level scaffolding pada pipeline RAG/orchestration, serta alur "
        "fitur forethought–performance–reflection (goal, diskusi terpantau, refleksi).",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
        first_line_indent=1.0,
        space_after=8,
    )
    add_para(doc, "2. Retrieval-Augmented Generation (RAG)", bold=True, space_after=4)
    add_para(
        doc,
        "Lewis et al. (2021) menetapkan paradigma RAG untuk mengondisikan generasi bahasa pada "
        "dokumen yang diambil. Di AI Engine, RAG diimplementasikan pada app/services/rag.py dengan "
        "kebijakan FETCH/NO_FETCH, retrieval Qdrant per course, reranker, grounding verifier, dan "
        "pembatasan konteks minggu (week_rag / weekContext di Core API).",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
        first_line_indent=1.0,
        space_after=8,
    )
    add_para(doc, "3. Self-Regulated Learning (Zimmerman)", bold=True, space_after=4)
    add_para(
        doc,
        "Kerangka Zimmerman (fase Forethought, Performance, Self-Reflection) menjadi acuan alur "
        "produk. Di kode, fase tersebut diwujudkan sebagai: (a) LearningGoal + validasi Bloom/AI "
        "goal_validator; (b) aktivitas diskusi, pre-read, usage/analytics, SRLClassifier berbasis "
        "pola pesan; (c) Reflection bertipe session/weekly beserta UI refleksi mahasiswa.",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
        first_line_indent=1.0,
        space_after=8,
    )
    add_para(doc, "4. Rekayasa kebutuhan sistem AI", bold=True, space_after=4)
    add_para(
        doc,
        "Hoy & Xu (2023) menyoroti ketidakpastian spesifikasi pada sistem berbasis AI. Kolabri "
        "menjawab ini dengan pemisahan layanan (scope boundaries Client App / Core API / AI Engine), "
        "konfigurasi provider AI di basis data (AiProvider), guardrail_policy per course, dan "
        "pengujian otomatis di ketiga layanan.",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
        first_line_indent=1.0,
        space_after=8,
    )
    add_para(doc, "5. Posisi sistem terhadap chatbot edukasi generik", bold=True, space_after=4)
    add_para(
        doc,
        "Berbeda dengan chatbot tanya-jawab tunggal, Kolabri menggabungkan (i) chat kelompok "
        "real-time multi-pengguna, (ii) chat AI pribadi, (iii) monitoring dosen, (iv) intervensi "
        "otomatis berbasis Logic Listener + eskalasi, dan (v) jejak data untuk analytics/process "
        "mining. Posisi ini adalah diferensiasi desain produk, bukan klaim superioritas empiris "
        "terhadap produk komersial tertentu.",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
        first_line_indent=1.0,
        space_after=8,
    )

    # ========== DETAILED DESIGN ==========
    add_heading(doc, "Detailed Design and Specification", level=1)
    add_para(
        doc,
        "Desain dan spesifikasi di bawah ini diturunkan dari implementasi repositori "
        "(Kolabri-client-app, Kolabri-core-api, Kolabri-ai-engine), bukan dari dokumen usang.",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
        first_line_indent=1.0,
        space_after=8,
    )

    add_para(doc, "1. Arsitektur sistem (tiga layanan)", bold=True, space_after=4)
    add_table(
        doc,
        ["Layanan", "Direktori / Port", "Teknologi utama (dari kode/deps)"],
        [
            [
                "Client App (BFF + UI)",
                "Kolabri-client-app\n:8000",
                "Laravel 12, PHP ^8.2, Inertia.js + React 19 + TypeScript, Vite, Tailwind, socket.io-client, Chart.js/Recharts",
            ],
            [
                "Core API",
                "Kolabri-core-api\n:3000",
                "Express 4, TypeScript, Prisma 6 + PostgreSQL, Mongoose/MongoDB, Socket.IO + Redis adapter, Zod, JWT",
            ],
            [
                "AI Engine",
                "Kolabri-ai-engine\n:8001",
                "FastAPI, Uvicorn, OpenAI-compatible LLM client, Qdrant, embedding lokal (paraphrase-multilingual-MiniLM-L12-v2), Motor/Mongo, Redis",
            ],
        ],
    )
    add_para(
        doc,
        "Infrastruktur pendukung di docker-compose.yml: PostgreSQL 16, MongoDB, Redis 7, Qdrant. "
        "Client App bertindak sebagai BFF: merender halaman Inertia dan mem-proxy sebagian request "
        "ke Core API; AI Engine dipanggil dari Core API (aiEngine.service) untuk RAG, validasi goal, "
        "intervensi, ringkasan, dan analytics NLP.",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
        first_line_indent=1.0,
        space_after=8,
    )

    add_para(doc, "2. Peran pengguna", bold=True, space_after=4)
    add_para(
        doc,
        "Enum UserRole di Prisma: student, lecturer, admin. Otorisasi di Core API (middleware role) "
        "dan di routes web Laravel (middleware role:student|lecturer|admin).",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
        first_line_indent=1.0,
        space_after=6,
    )
    add_bullets(
        doc,
        [
            "Student: join course/group, pre-read, goal, chat sesi, chat AI pribadi, refleksi, analytics personal.",
            "Lecturer: kelola course/group/session, knowledge base, materials/weeks, analytics, discussion health, attendance, eskalasi.",
            "Admin: user management, master data course, AI provider settings, usage stats, audit log.",
        ],
    )

    add_para(doc, "3. Spesifikasi fungsional (FR) — ringkas berdasar modul kode", bold=True, space_before=8, space_after=4)
    add_para(
        doc,
        "Tabel FR berikut merangkum kebutuhan fungsional yang terimplementasi di repositori "
        "(bukan daftar aspirasional). Setiap baris menunjuk modul/path sebagai bukti.",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
        first_line_indent=1.0,
        space_after=6,
    )
    add_table(
        doc,
        ["ID", "Modul", "Fungsionalitas", "Bukti implementasi (path/modul)"],
        [
            ["FR-AUTH-01", "Auth", "Login email/password, Google OAuth, JWT, verifikasi email, reset password", "auth routes/services; GoogleAuthController; token middleware"],
            ["FR-AUTH-02", "Auth", "RBAC student/lecturer/admin pada API dan routes Laravel", "Core API role middleware; web.php role:*"],
            ["FR-STU-01", "Student", "Course enrollment & group join", "course.service; group.service; CourseController"],
            ["FR-STU-02", "Student", "Pre-read completion gate sebelum join chat room", "SessionDiscussionPreReadCompletion; socket PRE_READ_REQUIRED"],
            ["FR-STU-03", "Student", "Learning goal bersama per sesi + validasi Bloom/AI", "goal.service; AI Engine goal_validator"],
            ["FR-STU-04", "Student", "Chat sesi kelompok real-time (send, pin, presence, history)", "socket/index.ts; ChatLog Mongo; ChatMessage PG"],
            ["FR-STU-05", "Student", "Chat AI pribadi (RAG + stream)", "AiChat/AiChatMessage; /chat/personal"],
            ["FR-STU-06", "Student", "Reflection session/weekly", "reflection.service; pages/student/reflections"],
            ["FR-STU-07", "Student", "Personal analytics dashboard", "student analytics pages/controllers"],
            ["FR-LEC-01", "Lecturer", "CRUD course, group, session discussion", "course/group/sessionDiscussion services + UI"],
            ["FR-LEC-02", "Lecturer", "Knowledge base & materials/weeks + PDF file serving", "knowledgeBase.service; LecturerMaterials*; CourseWeek*"],
            ["FR-LEC-03", "Lecturer", "Analytics overview/detail/comparison + radar", "pages/lecturer/analytics/**; dashboard.service"],
            ["FR-LEC-04", "Lecturer", "Discussion health monitoring", "discussion-health routes/UI"],
            ["FR-LEC-05", "Lecturer", "Attendance session & record", "AttendanceSession/Record; lecturer attendance"],
            ["FR-LEC-06", "Lecturer", "Eskalasi intervensi (nudge → probe → flag-lecturer → resolved)", "EscalationState; interventions.ts"],
            ["FR-ADM-01", "Admin", "User management & master data course", "pages/admin/user-management; master-data"],
            ["FR-ADM-02", "Admin", "AI provider settings, test connection, model discovery, fallback", "ai-provider routes; AI Engine admin.py"],
            ["FR-ADM-03", "Admin", "Usage stats (AiUsage) & audit log", "AiUsage; AuditLog; admin audit-log UI"],
            ["FR-AI-01", "AI", "RAG pipeline: FETCH/NO_FETCH, Qdrant, rerank, grounding, week cap", "rag.py; vector_store.py; week_rag.py"],
            ["FR-AI-02", "AI", "Guardrails input/output (injection, PII, Socratic, off-topic, toxicity)", "guardrails.py; socratic_filter"],
            ["FR-AI-03", "AI", "Logic Listener: silence, off-topic, inequity (Gini), quality", "logic_listener.py; interventionGate"],
            ["FR-AI-04", "AI", "Intervention NLG + orchestration /chat stream", "intervention.py; orchestration.py"],
            ["FR-AI-05", "AI", "Session summary, goal validation/feedback, NLP analytics", "summary routes; goals; nlp_analytics"],
            ["FR-RT-01", "Realtime", "Socket.IO + Redis adapter, presence, silence lock 10 menit", "socket/*; Redis adapter"],
            ["FR-LOG-01", "Logging", "ChatLog/ActivityLog runtime; AiUsage token/latency/cost", "Mongo ChatLog; Prisma AiUsage"],
            ["FR-LOG-02", "Logging", "Export job & process-mining helpers (XES/plan-vs-reality)", "ExportJob; xes_exporter.py; plan_vs_reality.py"],
        ],
    )

    add_para(doc, "4. Spesifikasi non-fungsional (NFR) yang terlihat di kode", bold=True, space_after=4)
    add_bullets(
        doc,
        [
            "Keamanan: helmet, rate limit (express-rate-limit / slowapi), JWT, role checks, XSS sanitization, injection detector, masking PII di pipeline AI.",
            "Real-time: Socket.IO dengan Redis adapter; presence room; silence lock di Redis.",
            "Ketahanan AI: circuit breaker service, fallback order provider (admin), degraded error handling LLM.",
            "Observabilitas: logging (Winston / structured logger AI), AiUsage (token, latency, cost), monitoring endpoints AI Engine.",
            "Kinerja chat: semantic cache / Redis di AI Engine; dashboard cache di Core API (modul dashboard-cache).",
            "Isolasi multi-course: filter course_id di Qdrant; otorisasi role di Core API/BFF.",
        ],
    )

    add_para(doc, "5. Model data utama", bold=True, space_before=8, space_after=4)
    add_para(
        doc,
        "Kolabri memakai polyglot storage dengan pemisahan tanggung jawab yang tegas: "
        "PostgreSQL adalah basis data utama untuk domain relasional; MongoDB dan Qdrant melengkapi "
        "kebutuhan log volume-tinggi dan vektor, bukan menggantikan domain PostgreSQL.",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
        first_line_indent=1.0,
        space_after=6,
    )
    add_para(doc, "PostgreSQL (Prisma di Core API) — sumber kebenaran domain:", bold=False, space_after=4)
    add_bullets(
        doc,
        [
            "User, Course, CourseStudent, Group, GroupMember",
            "SessionDiscussion, SessionDiscussionPreReadCompletion, ChatMessage",
            "LearningGoal, Reflection",
            "KnowledgeBase, CourseWeek, CourseMaterial, CourseWeekMaterial",
            "AiChat, AiChatMessage, AiProvider, AiUsage",
            "EscalationState, Notification, AuditLog, ExportJob",
            "AttendanceSession, AttendanceRecord",
        ],
    )
    add_para(doc, "MongoDB (Mongoose / Motor) — log runtime & analytics jejak:", bold=False, space_before=4, space_after=4)
    add_bullets(
        doc,
        [
            "ChatLog (pesan real-time + engagement analysis, pin, attachments)",
            "SilenceEvent, ActivityLog, mirror runtime eskalasi bila dipakai",
            "Log aktivitas/intervensi di AI Engine (mongodb_logger)",
        ],
    )
    add_para(doc, "Qdrant — vector store RAG:", bold=False, space_before=4, space_after=4)
    add_bullets(
        doc,
        [
            "Collection per course (course_{id}) dengan metadata filter course_id",
            "Dipakai AI Engine untuk retrieval; tidak menyimpan otorisasi atau entitas domain",
        ],
    )
    add_para(
        doc,
        "Client App (BFF) merender UI, mem-proxy ke Core API, dan menyimpan file PDF materi di "
        "disk storage. Metadata minggu/materi domain tetap berakar di PostgreSQL Core API; "
        "BFF dapat memegang indeks lokal untuk penyajian file yang diselaraskan lewat seed/"
        "ID deterministik agar pre-read dan stream PDF tetap konsisten.",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
        first_line_indent=1.0,
        space_after=8,
    )

    add_para(doc, "6. Alur perilaku utama", bold=True, space_after=4)
    add_para(doc, "6.1 Alur mahasiswa (diskusi terstruktur)", italic=True, space_after=4)
    add_numbered(
        doc,
        [
            "Login (email/password atau Google OAuth) → dashboard student.",
            "Masuk course → pilih group/session discussion yang terikat week.",
            "Pre-read materi minggu → complete pre-read (record SessionDiscussionPreReadCompletion).",
            "Buat/akses learning goal sesi (satu goal bersama; validasi Bloom + AI feedback).",
            "Join room Socket.IO; server menolak jika pre-read (dan gate terkait) belum terpenuhi.",
            "Kirim pesan; Core API menyimpan log, memanggil AI Engine orchestration untuk respons RAG + kemungkinan intervensi.",
            "Tutup sesi → generate/simpan summary; submit reflection (session dan/atau weekly).",
        ],
    )
    add_para(doc, "6.2 Pipeline AI orchestration (ringkas)", italic=True, space_before=6, space_after=4)
    add_numbered(
        doc,
        [
            "Analisis NLP pesan (engagement/HOT/lexical bila dijalankan).",
            "RAG query/stream: policy FETCH/NO_FETCH → retrieval Qdrant → rerank → generasi LLM → sitasi/grounding metadata.",
            "Penerapan guardrails pada input/output; filter Socratic bila jawaban terlalu langsung.",
            "Logging aktivitas ke Mongo.",
            "Logic Listener / intervention service: off-topic, silence, participation inequity, quality issues.",
            "Bila perlu, generate pesan intervensi; eskalasi bertahap hingga flag-lecturer.",
        ],
    )
    add_para(doc, "6.3 Ambang operasional yang terdefinisi di kode", italic=True, space_before=6, space_after=4)
    add_bullets(
        doc,
        [
            "Silence timeout socket gate: 10 menit (SILENCE_TIMEOUT_MS = 10 * 60 * 1000); SILENCE_THRESHOLD_MINUTES default 10 di AI Engine.",
            "Cooldown intervensi: 3 menit; quality check setelah kelipatan 5 pesan (MESSAGES_BEFORE_CHECK).",
            "Logic Listener: similarity off-topic default 0.6; consecutive off-topic default 3; inequity (Gini) threshold default 0.6.",
            "Eskalasi default: nudgeAfterMs 5 menit, probeAfterMs 10 menit (dapat dioverride aiEscalationConfig per course).",
            "Tahapan eskalasi: new → nudge → probe-blocker → flag-lecturer → resolved.",
        ],
    )

    add_para(doc, "7. Antarmuka (UI) per peran — halaman yang ada di resources/js/pages", bold=True, space_before=8, space_after=4)
    add_bullets(
        doc,
        [
            "Student: dashboard, courses (index/show/attendance), groups/show, pre-read/show, goals/create, chat room (+ pin/search/edit), ai-chat, reflections, profile.",
            "Lecturer: dashboard, courses (create/index/show), groups, analytics (overview/detail/comparison/show), RadarChartPage; materials & course-weeks dikelola lewat controller/UI lecturer.",
            "Admin: dashboard, user-management, master-data, ai-settings, audit-log.",
            "Auth/settings: login/register/forgot-password/reset-password/verify-email, settings (profile, security, appearance, notifications).",
        ],
    )

    # ========== DEVELOPMENT ==========
    add_heading(doc, "Development", level=1)
    add_para(doc, "1. Lingkungan pengembangan", bold=True, space_after=4)
    add_para(doc, "Perangkat lunak yang dibutuhkan (sesuai package/runtime proyek):", space_after=4)
    add_bullets(
        doc,
        [
            "Node.js (Core API & tooling client), npm",
            "PHP 8.2+ dan Composer (Laravel client)",
            "Python 3.11+ (AI Engine; pyproject requires-python >=3.11)",
            "PostgreSQL, MongoDB, Redis, Qdrant",
            "Git; opsional Docker Compose untuk stack infrastruktur/services",
        ],
    )
    add_para(doc, "Variabel lingkungan inti (dari .env.example):", space_before=6, space_after=4)
    add_bullets(
        doc,
        [
            "Core API: PORT=3000, DATABASE_URL (PostgreSQL), MONGODB_URL, AI_ENGINE_URL=http://localhost:8001, JWT_SECRET, Redis",
            "AI Engine: PORT=8001, OPENAI_BASE_URL / OPENAI_API_KEY / OPENAI_MODEL, QDRANT_URL, embedding model, flag UNIFIED_PROVIDER_*",
            "Client App: APP_URL=http://localhost:8000, konfigurasi DB BFF, URL Core API, secrets OAuth Google bila dipakai",
        ],
    )
    add_para(doc, "Menjalankan lokal (pola dev.sh / manual):", space_before=6, space_after=4)
    add_numbered(
        doc,
        [
            "Start infrastruktur (postgres, mongo, redis, qdrant).",
            "Start AI Engine: uvicorn/main.py di :8001.",
            "Start Core API: npm run dev di :3000.",
            "Start Client App: php artisan serve (+ vite/build) di :8000.",
            "Seed demo: scripts/seed-demo.sh atau npm run db:reset:demo-data + MaterialsDemoSeeder.",
        ],
    )

    add_para(doc, "2. Proses pengembangan", bold=True, space_before=8, space_after=4)
    add_para(
        doc,
        "Pengembangan mengikuti pemisahan domain antar repositori layanan dengan kontrak HTTP/JSON "
        "dan Socket.IO. Perubahan AI tidak diletakkan di UI; Client App mem-proxy ke Core API; "
        "Core API memanggil AI Engine. Fitur pedagogis dilengkapi model data (goal, pre-read, "
        "reflection, escalation) agar alur dapat diverifikasi end-to-end. Bukti proses terlihat dari:",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
        first_line_indent=1.0,
        space_after=6,
    )
    add_bullets(
        doc,
        [
            "Struktur monorepo layanan: Kolabri-core-api/, Kolabri-ai-engine/, Kolabri-client-app/",
            "Suite pengujian: ratusan berkas test di core-api (*.test.ts), ai-engine (tests/), dan client (PHPUnit + Vitest-style tests di resources/js)",
            "Skrip operasional: dev.sh, docker-compose.yml, scripts/seed-demo.sh",
            "Konfigurasi provider AI di DB + admin UI (bukan hardcode tunggal di frontend)",
            "ADR dan boundary docs di docs/architecture sebagai pelengkap (bukan sumber kebenaran tunggal)",
        ],
    )

    add_para(doc, "3. Pemetaan peran ke artefak kode (bukti development)", bold=True, space_before=8, space_after=4)
    add_para(
        doc,
        "Tabel berikut memetakan anggota tim ke area domain dan artefak kode yang menjadi "
        "jejak kerja teknis (bukan klaim ownership eksklusif git blame).",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
        first_line_indent=1.0,
        space_after=6,
    )
    add_table(
        doc,
        ["Anggota / peran", "Area domain", "Contoh artefak di repositori"],
        [
            [
                "Soraya — System Analyst",
                "Spesifikasi & alur",
                "Kontrak role routes web; FR/NFR di laporan; gate pre-read/goal di BFF+socket; docs/architecture boundaries",
            ],
            [
                "Irham — Backend Chat & User",
                "Core API domain + realtime",
                "Kolabri-core-api/src/socket/*; group.service; auth.service; sessionDiscussion.service; interventionGate.ts; Prisma schema domain",
            ],
            [
                "Hashfi — Backend AI & Analytics",
                "AI Engine + integrasi",
                "Kolabri-ai-engine/app/services/rag.py; orchestration.py; logic_listener.py; guardrails; nlp_analytics; plan_vs_reality; xes_exporter; core-api aiEngine.service",
            ],
            [
                "Rizky — Frontend Chat & Group",
                "UI mahasiswa diskusi",
                "pages/student/chat/**; pre-read; goals; groups; ai-chat; socket.io-client hooks & message components",
            ],
            [
                "Akbar — Frontend Dashboard",
                "UI dosen/admin analytics",
                "pages/lecturer/analytics/**; dashboard; RadarChartPage; pages/admin/*; Chart.js/Recharts widgets",
            ],
            [
                "Arvin — QA",
                "Pengujian multi-layanan",
                "core-api *.test.ts; ai-engine/tests/**; client tests/ Feature+JS; smoke/UAT di docs/testing & docs/questionnaires",
            ],
        ],
    )

    # ========== EVALUATION ==========
    add_heading(doc, "Evaluation Process, Result, and Analysis", level=1)
    add_para(
        doc,
        "Bagian evaluasi membatasi diri pada mekanisme dan artefak pengujian yang ada di repositori. "
        "Angka performa LLM atau skor pedagogis lapangan tidak diisi di sini kecuali diukur ulang "
        "secara eksplisit; laporan ini menghindari klaim yang tidak tertambat pada kode/test.",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
        first_line_indent=1.0,
        space_after=8,
    )

    add_para(doc, "1. Skenario evaluasi", bold=True, space_after=4)
    add_bullets(
        doc,
        [
            "Pengujian unit/integration Core API (Vitest/Jest-style *.test.ts pada services, socket, routes).",
            "Pengujian AI Engine (pytest di tests/: RAG, guardrails, orchestration, interventions, analytics, admin routes, dsb.).",
            "Pengujian Client App (tests/ Laravel + unit frontend chat/summary/drag-drop).",
            "Skenario fungsional end-to-end manual/smoke: login multi-role, course/group, pre-read→goal→chat→close→reflection, analytics dosen, admin AI settings.",
            "UAT/kuesioner tersedia sebagai instrumen di docs/questionnaires (dosen & mahasiswa) untuk evaluasi pengguna; hasil lapangan terpisah dari verifikasi perangkat lunak.",
        ],
    )

    add_para(doc, "2. Pelaksanaan evaluasi (proses)", bold=True, space_before=6, space_after=4)
    add_numbered(
        doc,
        [
            "Menjalankan suite test per layanan (npm test di core-api; pytest di ai-engine; php artisan test / npm test di client sesuai konfigurasi proyek).",
            "Memverifikasi alur gerbang pedagogis (pre-read & goal) baik di BFF maupun socket (kode error PRE_READ_REQUIRED).",
            "Memverifikasi cabang intervensi: silence lock, quality thresholds (HOT/cognitive/lexical), off-topic listener.",
            "Memverifikasi integrasi AI: ingest knowledge base → vector status; query RAG; personal chat stream; provider test admin.",
            "Memverifikasi peran: endpoint admin/lecturer/student menolak role yang tidak berwenang.",
        ],
    )

    add_para(doc, "3. Hasil (temuan implementasi — factual dari audit kode)", bold=True, space_before=6, space_after=4)
    add_bullets(
        doc,
        [
            "Empat masalah proposal (partisipasi tidak merata, monitoring dosen, off-topic, ringkasan) memiliki path implementasi: analytics partisipasi/Gini/listener; dashboard & discussion health; off-topic detection; summary pada SessionDiscussion dan endpoint summary/regenerate.",
            "Dua mode chat tersedia: session group (SessionDiscussion + socket + orchestration) dan personal (AiChat).",
            "Siklus Zimmerman terpetakan ke fitur: goal (forethought), aktivitas diskusi/pre-read/analytics (performance), reflection (self-reflection); ditambah SRLClassifier berbasis pola teks.",
            "Empat guardrail pedagogis/safety tercakup di modul: no direct answers (Socratic), no complete solution / academic dishonesty, injection prevention, stay-in-context/off-topic + grounding metadata.",
            "Research-oriented logging: AiUsage, AuditLog, ChatLog/ActivityLog, ExportJob, ekspor XES, plan-vs-reality analyzer.",
            "Jumlah berkas uji yang teramati di repositori (order-of-magnitude, dapat berubah seiring commit): ~230 berkas *.test.ts di core-api, ~140 modul test_*.py di ai-engine, ~50 test PHP/JS di client — menandakan evaluasi perangkat lunak diotomatisasi, bukan hanya demo manual.",
            "Bukti reproduksi AI Engine (Juli 2026, docs/evidence/bab4/): suite pytest+coverage dijalankan; Locust 20 users engagement ~65 RPS (0 fail); HTTP engagement ~115 RPS; RAG rerank menaikkan P@3 pada gold 20 kueri; detail angka di REPRODUCTION_REPORT_20260721.md. Angka ini validasi teknis, bukan efektivitas pedagogis lapangan.",
        ],
    )

    add_para(doc, "4. Analisis", bold=True, space_before=6, space_after=4)
    add_para(
        doc,
        "Mengapa solusi menjawab problem definition: pemisahan layanan memungkinkan Core API menjaga "
        "integritas domain (user, course, session, socket) sementara AI Engine mengisolasi beban "
        "compute (embedding, retrieval, LLM). Gerbang pre-read dan goal memaksa fase forethought "
        "sebelum performance di ruang chat. Logic Listener dan eskalasi menutup gap monitoring dosen "
        "tanpa mengharuskan kehadiran manual di setiap kelompok. RAG + guardrails menempatkan "
        "respons AI pada materi kursus dan kebijakan akademik, selaras dengan risiko halusinasi "
        "yang diuraikan di latar belakang.",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
        first_line_indent=1.0,
        space_after=8,
    )
    add_para(
        doc,
        "Batasan interpretasi hasil: lulusnya test otomatis membuktikan kebenaran fungsional/"
        "regresi teknis pada skenario yang ditulis, bukan otomatis membuktikan peningkatan "
        "capaian belajar mahasiswa di kelas nyata. Validasi pedagogis memerlukan studi pengguna/"
        "eksperimen terpisah (instrumen UAT/kuesioner disiapkan, tetapi tidak diganti oleh unit test).",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
        first_line_indent=1.0,
        space_after=8,
    )

    # ========== CHALLENGES ==========
    add_heading(doc, "Challenges", level=1)
    add_bullets(
        doc,
        [
            "Pemisahan store polyglot: domain utama di PostgreSQL; volume chat/event di MongoDB; embedding di Qdrant. Konsistensi ID week/material (deterministik) diperlukan agar pre-read BFF, Core API, dan seed tetap selaras saat PDF disajikan dari storage client.",
            "Konsistensi gate pedagogis: pre-read/goal harus ditegakkan di BFF (HTTP page) dan di socket join agar tidak bisa dibypass lewat klien WebSocket.",
            "Ambiguitas spesifikasi AI: ambang silence/off-topic/inequity dan kebijakan guardrail perlu dikonfigurasi (settings + per-course JSON) agar tidak “keras-tertanam” tanpa konteks kelas.",
            "Orkestrasi multi-komponen per pesan: RAG + logging + intervensi + guardrails berimplikasi pada latensi dan keharusan streaming (/chat/stream, personal stream).",
            "Provider LLM heterogen: abstraksi OpenAI-compatible + AiProvider DB + fallback order; flag UNIFIED_PROVIDER_* mengatur sumber konfigurasi (env vs DB).",
            "Keandalan real-time: presence, silence lock Redis, cooldown intervensi, dan pencegahan double-nudge pada banyak instance.",
            "Keseimbangan scaffolding: mencegah jawaban langsung (Socratic) tanpa membuat asisten tidak berguna; diimplementasikan sebagai filter/rewrite, bukan sekadar system prompt statis.",
            "Evaluasi: membedakan verifikasi software (test suite + repro performa) dari validasi ilmiah pedagogis agar klaim laporan tidak melebihi bukti.",
        ],
    )

    # ========== CONCLUSION ==========
    add_heading(doc, "Conclusion", level=1)
    add_para(
        doc,
        "Kolabri telah diimplementasikan sebagai platform web chatbot kolaboratif yang menyatukan "
        "diskusi kelompok real-time, chat AI pribadi, monitoring dosen, intervensi otomatis, RAG "
        "berbasis materi, dan jejak data analytics. Arsitektur tiga layanan (Client App, Core API, "
        "AI Engine) merealisasikan pemisahan tanggung jawab yang selaras dengan problem definition: "
        "regulasi belajar terfasilitasi lewat pre-read–goal–diskusi–refleksi; dinamika kelompok "
        "dipantau Logic Listener dan eskalasi; akurasi/keamanan AI ditopang RAG dan guardrails; "
        "data terstruktur tersedia untuk monitoring dan analisis lanjutan.",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
        first_line_indent=1.0,
        space_after=8,
    )
    add_para(
        doc,
        "Sub-permasalahan per peran terjawab pada level rekayasa perangkat lunak: spesifikasi alur "
        "dan kontrak layanan; backend chat/user dengan Socket.IO dan model relasional/dokumen; "
        "backend AI dengan pipeline orchestration; frontend chat dan dashboard; serta fondasi QA "
        "berupa suite pengujian multi-layanan. Lingkup laporan ini tidak mengklaim evaluasi "
        "efektivitas pembelajaran di lapangan melampaui bukti yang ada di repositori.",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
        first_line_indent=1.0,
        space_after=8,
    )

    # ========== ANNEX ==========
    add_heading(doc, "Annex (e.g. user manual)", level=1)
    add_para(doc, "A. Kredensial demo (dari seed proyek; dapat berubah sesuai environment)", bold=True, space_after=4)
    add_bullets(
        doc,
        [
            "Dosen contoh (seed Agents.md): budi.santoso@univ.ac.id / password123; siti.rahayu@univ.ac.id / password123",
            "Mahasiswa contoh: andi.pratama@student.ac.id / password123",
            "Alternatif demo guide historis: lecturer@kolabri.edu / student1@kolabri.edu — gunakan yang aktif di seed environment Anda",
        ],
    )
    add_para(doc, "B. Manual ringkas per peran", bold=True, space_before=8, space_after=4)
    add_para(doc, "Mahasiswa", italic=True, space_after=4)
    add_numbered(
        doc,
        [
            "Buka http://localhost:8000 → login.",
            "Pilih mata kuliah → buka sesi diskusi minggu terkait.",
            "Selesaikan pre-read → tetapkan goal → masuk chat room.",
            "Diskusikan dengan rekan; gunakan AI sesuai kebijakan kursus; tutup sesi dan isi refleksi.",
            "Opsional: gunakan menu AI Chat untuk konsultasi 1:1 berbasis materi.",
        ],
    )
    add_para(doc, "Dosen", italic=True, space_before=6, space_after=4)
    add_numbered(
        doc,
        [
            "Login sebagai lecturer → dashboard.",
            "Kelola course, group, session discussion, unggah knowledge base.",
            "Pantau analytics dan discussion health; tindak lanjuti eskalasi bila ada.",
            "Tinjau ringkasan sesi dan partisipasi mahasiswa.",
        ],
    )
    add_para(doc, "Admin", italic=True, space_before=6, space_after=4)
    add_numbered(
        doc,
        [
            "Login admin → user management & master data.",
            "Atur AI Settings (provider, activate, fallback, test connection, daftar model).",
            "Pantau usage stats (AiUsage) dan audit log.",
        ],
    )
    add_para(doc, "C. Endpoint AI Engine (kelompok modul)", bold=True, space_before=8, space_after=4)
    add_bullets(
        doc,
        [
            "Health & monitoring: health, metrics, circuit-breakers, reranker health",
            "Documents: ingest / batch ingest / delete collection data",
            "Chat & RAG: /ask, personal chat (+ stream), reading recommendations",
            "Orchestration: /chat, /chat/stream",
            "Interventions: analyze, summary, prompt",
            "Goals: validasi/feedback goal",
            "Analytics & groups: metrik grup, efficiency, plan/analytics routes",
            "Discussion direction: classify-relevance, session-summary",
            "Admin: test-provider, providers/{provider}/models",
        ],
    )
    add_para(doc, "D. Struktur repositori", bold=True, space_before=8, space_after=4)
    add_bullets(
        doc,
        [
            "Kolabri-client-app/ — UI + BFF Laravel",
            "Kolabri-core-api/ — domain API + socket + Prisma schema",
            "Kolabri-ai-engine/ — RAG, LLM, analytics, interventions",
            "docs/ — arsitektur, testing, kuesioner (referensi pendukung)",
            "D3/ — laporan capstone final (dokumen ini) + skrip extract/compile",
            "dev.sh, docker-compose.yml, scripts/seed-demo.sh — operasional lokal",
        ],
    )

    # ========== REFERENCES ==========
    add_heading(doc, "Referensi", level=1)
    refs = [
        "Hoy, Z., & Xu, M. (2023). Challenges of AI system requirements engineering. (acuancut dalam perumusan background D3).",
        "Lewis, P., Perez, E., Piktus, A., Petroni, F., Karpukhin, V., Goyal, N., … Kiela, D. (2021). Retrieval-Augmented Generation for Knowledge-Intensive NLP Tasks. Advances in Neural Information Processing Systems.",
        "Martha, A. S. D., et al. (2023). Adaptive scaffolding in pedagogical agents for self- and co-regulated learning. (acuancut dalam background D3).",
        "Zimmerman, B. J. (2002). Becoming a self-regulated learner: An overview. Theory Into Practice, 41(2), 64–70.",
        "Repositori implementasi Kolabri: Kolabri-core-api, Kolabri-ai-engine, Kolabri-client-app (source of truth desain dan spesifikasi laporan ini).",
        "Prisma schema: Kolabri-core-api/prisma/schema.prisma.",
        "AI Engine entrypoint & config: Kolabri-ai-engine/main.py; app/core/config.py.",
        "Socket real-time & gates: Kolabri-core-api/src/socket/index.ts; interventionGate.ts; engagement.ts.",
        "RAG & Logic Listener: Kolabri-ai-engine/app/services/rag.py; logic_listener.py; orchestration.py; guardrails.py.",
    ]
    for i, r in enumerate(refs, 1):
        add_para(doc, f"[{i}] {r}", align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=4, size=11)

    doc.save(str(OUT))
    print(f"Wrote {OUT}")
    return OUT


if __name__ == "__main__":
    build()
